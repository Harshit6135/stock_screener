"""Bounded extraction of contract-note and funds-statement charge rows."""

import csv
import io
import re
from datetime import UTC, datetime
from decimal import ROUND_DOWN, Decimal, InvalidOperation
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from src.domains.portfolio_accounting.tradebook import canonical_symbol
from src.platform_kernel import DomainValidationError

COMPONENT_LABELS = {
    "brokerage": "Brokerage",
    "exchange": "Exchange transaction charges",
    "clearing": "Clearing charges",
    "sebi": "SEBI charges",
    "gst": "GST",
    "stt": "STT",
    "stamp_duty": "Stamp duty",
    "dp": "DP charges",
    "other": "Other / unitemized charges",
    "recorded_fee": "Previously recorded fee",
}
ALIASES = {
    "date": {"date", "tradedate", "transactiondate", "postingdate"},
    "symbol": {"symbol", "tradingsymbol", "security", "scrip"},
    "isin": {"isin"},
    "side": {"side", "tradetype", "transactiontype", "buysell"},
    "units": {"quantity", "qty", "units"},
    "price": {"price", "tradeprice", "grossrate"},
    "trade_id": {"tradeid", "tradeno", "tradenumber"},
    "exchange": {"exchange"},
    "description": {"description", "particulars", "narration"},
    "debit": {"debit", "debits", "debitamount"},
    "credit": {"credit", "credits", "creditamount"},
    "brokerage": {"brokerage", "brokeragecharges"},
    "exchange_charge": {"exchangecharges", "exchangetransactioncharges", "transactioncharges"},
    "sebi": {"sebi", "sebicharges", "sebiturnovercharges"},
    "gst": {"gst", "gstcharges", "igst"},
    "cgst": {"cgst"},
    "sgst": {"sgst"},
    "stt": {"stt", "securitiestransactiontax"},
    "stamp_duty": {"stampduty", "stampcharges"},
    "dp": {"dp", "dpcharges", "depositorycharges"},
    "total": {"charges", "totalcharges", "fee", "fees"},
}


def amount(value):
    try:
        number = Decimal(str(value or "0").replace(",", "").replace("₹", "").strip())
    except InvalidOperation as exc:
        raise DomainValidationError("invalid charge amount") from exc
    if not number.is_finite() or number < 0 or number > Decimal(10000000):
        raise DomainValidationError("charge amounts must be nonnegative and finite")
    return number


def iso_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d %b %Y"):
        try:
            return (
                datetime.strptime(str(value).strip()[:11], fmt)
                .replace(tzinfo=UTC)
                .date()
                .isoformat()
            )
        except ValueError:
            continue
    raise DomainValidationError("document row has an invalid date")


def allocate(total, weights):
    """Conserve the source total, including shares attributable to outside trades."""
    total = amount(total).quantize(Decimal("0.01"))
    denominator = sum(weights, Decimal(0))
    if not denominator:
        return [Decimal(0)] * len(weights)
    exact = [total * weight / denominator for weight in weights]
    parts = [value.quantize(Decimal("0.01"), rounding=ROUND_DOWN) for value in exact]
    pennies = int((total - sum(parts)) * 100)
    order = sorted(range(len(weights)), key=lambda i: exact[i] - parts[i], reverse=True)
    for index in order[:pennies]:
        parts[index] += Decimal("0.01")
    return parts


def _tables(raw, extension):
    if extension == ".csv":
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DomainValidationError("CSV must use UTF-8 encoding") from exc
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [list(csv.reader(io.StringIO(text), dialect))]
    with ZipFile(io.BytesIO(raw)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
            raise DomainValidationError("expanded spreadsheet exceeds 50 MB")
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        tables = []
        for sheet in workbook.worksheets[:10]:
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= 20000:
                    raise DomainValidationError("spreadsheet exceeds 20,000 rows per sheet")
                rows.append(list(row[:40]))
            tables.append(rows)
        return tables
    finally:
        workbook.close()


def _parse_tables(tables, kind):
    result, warnings = [], []
    for table in tables:
        columns, start = None, None
        for index, row in enumerate(table[:60]):
            names = [re.sub(r"[^a-z0-9]", "", str(cell or "").lower()) for cell in row]
            found = {
                key: next((i for i, name in enumerate(names) if name in aliases), None)
                for key, aliases in ALIASES.items()
            }
            if found["date"] is not None and (
                found["side"] is not None
                or found["description"] is not None
                or found["dp"] is not None
            ):
                columns, start = found, index + 1
                break
        if columns is None:
            continue
        for line, row in enumerate(table[start:], start + 1):
            if not any(cell is not None and str(cell).strip() for cell in row):
                continue

            def cell(key, columns=columns, row=row):
                index = columns[key]
                return row[index] if index is not None and index < len(row) else ""

            description = str(cell("description") or "")
            statement = kind == "statement" or columns["side"] is None
            if statement and not (
                "dp" in description.lower()
                or "depositor" in description.lower()
                or columns["dp"] is not None
            ):
                continue
            try:
                if statement and cell("credit") and amount(cell("credit")):
                    warnings.append(
                        f"Row {line}: DP credit/refund requires review; it was not imported as a debit."
                    )
                    continue
                day = iso_date(cell("date"))
                symbol = canonical_symbol(str(cell("symbol") or ""))
                isin = str(cell("isin") or "").strip()
                if not isin:
                    found_isin = re.search(r"\bIN[A-Z0-9]{10}\b", description.upper())
                    isin = found_isin.group() if found_isin else ""
                if statement:
                    fee = amount(cell("dp") or cell("debit") or cell("total"))
                    if not fee:
                        continue
                    result.append(
                        {
                            "date": day,
                            "symbol": symbol,
                            "isin": isin,
                            "side": "SELL",
                            "units": 0,
                            "price": "0",
                            "trade_id": "",
                            "exchange": "NSE",
                            "components": {"dp": str(fee)},
                            "group": "dp",
                            "estimated": False,
                            "line": line,
                            "description": description[:200],
                        }
                    )
                    continue
                side = str(cell("side")).upper().strip()
                side = {"B": "BUY", "S": "SELL"}.get(side, side)
                units_value = amount(cell("units"))
                if (
                    side not in {"BUY", "SELL"}
                    or units_value <= 0
                    or units_value != int(units_value)
                ):
                    raise DomainValidationError("invalid trade side or quantity")
                components = {}
                for key in ("brokerage", "sebi", "gst", "stt", "stamp_duty", "dp"):
                    if columns[key] is not None:
                        components[key] = str(amount(cell(key)))
                if columns["exchange_charge"] is not None:
                    components["exchange"] = str(amount(cell("exchange_charge")))
                if columns["cgst"] is not None or columns["sgst"] is not None:
                    split_gst = amount(cell("cgst")) + amount(cell("sgst"))
                    declared_gst = amount(components.get("gst", "0"))
                    if declared_gst and split_gst and declared_gst != split_gst:
                        raise DomainValidationError("GST total disagrees with CGST and SGST")
                    components["gst"] = str(declared_gst or split_gst)
                if columns["total"] is not None:
                    total = amount(cell("total"))
                    known = sum((amount(v) for v in components.values()), Decimal(0))
                    if known > total:
                        raise DomainValidationError("component charges exceed the declared total")
                    if not components or total > known:
                        components["other"] = str(total - known)
                if not components:
                    raise DomainValidationError("no charge columns found")
                price = amount(cell("price"))
                if price <= 0:
                    raise DomainValidationError("trade price must be positive")
                result.append(
                    {
                        "date": day,
                        "symbol": symbol,
                        "isin": isin,
                        "side": side,
                        "units": int(units_value),
                        "price": str(price),
                        "trade_id": str(cell("trade_id") or "").strip(),
                        "exchange": str(cell("exchange") or "NSE").upper(),
                        "components": components,
                        "group": "contract",
                        "estimated": False,
                        "line": line,
                    }
                )
            except DomainValidationError as exc:
                warnings.append(f"Row {line}: {exc}")
    if not result:
        warnings.append(
            "No charge rows recognized. Use the CSV template or review the document format."
        )
    return result, warnings


def _pdf_text(raw, password):
    try:
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted and not reader.decrypt(password or ""):
            raise DomainValidationError("PDF password is missing or incorrect")
        if len(reader.pages) > 100:
            raise DomainValidationError("PDF exceeds 100 pages")
        pages = [page.extract_text(extraction_mode="layout") or "" for page in reader.pages]
        if sum(map(len, pages)) > 2_000_000:
            raise DomainValidationError("extracted document is too large")
        return "\n".join(pages)
    except (PdfReadError, ValueError, NotImplementedError) as exc:
        raise DomainValidationError("PDF could not be read; check its password and format") from exc


def _parse_pdf(text, kind):
    # Normalized table exports embedded in PDFs are accepted as well.
    if "symbol,date,side,quantity,price" in text.lower():
        start = text.lower().index("symbol,date,side,quantity,price")
        return _parse_tables([list(csv.reader(io.StringIO(text[start:])))], kind)
    if kind == "statement":
        rows = []
        for line, text_line in enumerate(text.splitlines(), 1):
            if not re.search(r"\bDP\s*(?:charges|charge)\b", text_line, re.IGNORECASE):
                continue
            day = re.search(r"\b(?:\d{4}-\d{2}-\d{2}|\d{2}[-/]\d{2}[-/]\d{4})\b", text_line)
            isin = re.search(r"\bIN[A-Z0-9]{10}\b", text_line.upper())
            # Only an explicit debit is accepted from free-form PDF statements.
            debit = re.search(r"(?:debit|dr)\s*[:₹]?\s*([\d,]+\.\d{2})", text_line, re.IGNORECASE)
            if day and debit:
                rows.append(
                    {
                        "date": iso_date(day.group()),
                        "symbol": "",
                        "isin": isin.group() if isin else "",
                        "side": "SELL",
                        "units": 0,
                        "price": "0",
                        "trade_id": "",
                        "exchange": "NSE",
                        "components": {"dp": str(amount(debit.group(1)))},
                        "group": "dp",
                        "estimated": False,
                        "line": line,
                        "description": "DP statement debit",
                    }
                )
        return rows, [] if rows else [
            "PDF statement layout not recognized. Import the funds statement CSV/XLSX export instead."
        ]
    dates = re.findall(
        r"(?:trade\s*date|date\s*of\s*trade)\s*[:\-]?\s*(\d{4}-\d{2}-\d{2}|\d{2}[-/]\d{2}[-/]\d{4})",
        text,
        re.IGNORECASE,
    )
    days = {iso_date(day) for day in dates}
    if len(days) != 1:
        return [], [
            "Contract note needs one identifiable trade date; use one note per file or the CSV template."
        ]
    day = next(iter(days))
    if "Net Rate per" in text and "Net Obligation for ISIN" in text:
        return _zerodha_annexure(text, day)
    pattern = re.compile(
        r"^\s*\d{5,}\s+\d{2}:\d{2}:\d{2}\s+(\d+)\s+\d{2}:\d{2}:\d{2}\s+(.+?)\s+([BS])\s+(NSE|BSE)\s+(.+)$"
    )
    rows, totals, warnings, seen = [], {}, [], set()
    gst_parts = {}
    for line, text_line in enumerate(text.splitlines(), 1):
        match = pattern.match(text_line)
        if match:
            trade_id, symbol, side, exchange, rest = match.groups()
            if (exchange, trade_id) in seen:
                continue
            numbers = re.findall(r"-?[\d,]+(?:\.\d+)?", rest)
            if len(numbers) < 3:
                warnings.append(
                    f"Trade row {line} was not recognized; aggregate charges were not allocated."
                )
                continue
            units = amount(numbers[0].lstrip("-"))
            # Zerodha's legacy equity row has '-' for foreign-currency rate.
            price = amount(numbers[1])
            if units <= 0 or units != int(units) or price <= 0:
                warnings.append(f"Trade row {line} has invalid quantity or price.")
                continue
            brokerage = amount(numbers[2]) * units
            rows.append(
                {
                    "date": day,
                    "symbol": canonical_symbol(symbol.split("(")[0].strip()),
                    "isin": "",
                    "side": "BUY" if side == "B" else "SELL",
                    "units": int(units),
                    "price": str(price),
                    "trade_id": trade_id,
                    "exchange": exchange,
                    "components": {"brokerage": str(brokerage)},
                    "group": "contract",
                    "estimated": False,
                    "line": line,
                }
            )
            seen.add((exchange, trade_id))
            continue
        if re.match(r"^\s*\d{5,}\s+\d{2}:\d{2}:\d{2}", text_line):
            warnings.append(
                f"Trade row {line} has an unsupported layout; aggregate charges were not allocated."
            )
        lower = text_line.lower().strip()
        labels = {
            "exchange": r"(?:exchange|transaction)\s*(?:transaction\s*)?charges",
            "sebi": r"sebi\s*(?:turnover\s*)?(?:fees?|charges)",
            "stt": r"(?:securities transaction tax|stt)\b",
            "stamp_duty": r"stamp\s*duty",
            "gst": r"(?:igst|cgst|sgst|gst)\b",
            "brokerage": r"(?:total\s*)?brokerage\b",
        }
        for key, label in labels.items():
            if re.match(label, lower):
                numbers = re.findall(r"\(?([\d,]+\.\d{2,})\)?", text_line)
                if numbers:
                    values = [amount(n) for n in numbers]
                    # Some notes have a final Total column beside segment values.
                    value = (
                        values[-1]
                        if len(values) > 1 and values[-1] == sum(values[:-1])
                        else sum(values, Decimal(0))
                    )
                    total_key = (
                        re.match(r"(?:igst|cgst|sgst|gst)", lower).group() if key == "gst" else key
                    )
                    destination = gst_parts if key == "gst" else totals
                    if total_key in destination and destination[total_key] != value:
                        warnings.append(
                            f"Conflicting {total_key} summaries; use a normalized export."
                        )
                    destination[total_key] = value
                break
    if gst_parts:
        split_gst = sum((v for k, v in gst_parts.items() if k != "gst"), Decimal(0))
        if gst_parts.get("gst") and split_gst and gst_parts["gst"] != split_gst:
            warnings.append("GST total disagrees with its components; use a normalized export.")
        totals["gst"] = gst_parts.get("gst") or split_gst
    if not rows:
        return [], [
            "Contract-note trade table not recognized. Use the CSV template or provide this format for parser support."
        ]
    if warnings:
        return [], warnings  # Never allocate daily totals over an incomplete trade table.
    if not totals:
        return [], [
            "No contract-note charge summary recognized; brokerage alone is not a complete charge import."
        ]
    for component, total in totals.items():
        if (
            component == "brokerage"
            and sum((amount(r["components"]["brokerage"]) for r in rows), Decimal(0)) == total
        ):
            continue
        weights = [
            amount(r["price"]) * r["units"]
            if component != "stamp_duty" or r["side"] == "BUY"
            else Decimal(0)
            for r in rows
        ]
        for row, share in zip(rows, allocate(total, weights), strict=True):
            row["components"][component] = str(share)
            row["estimated"] = True
    warnings.append(
        "Daily summary charges are turnover allocations across every recognized trade, including outside trades; these are estimates."
    )
    return rows, warnings


def _zerodha_annexure(text, day):
    """Read equity annexures and the cash segment of newer Zerodha notes."""
    lines = text.replace("−", "-").splitlines()
    trade_pattern = re.compile(
        r"^\s*\d{5,}\s+\d{2}:\d{2}:\d{2}\s+(\d+)\s+\d{2}:\d{2}:\d{2}\s+"
        r"(.+?)/(?P<isin>IN[A-Z0-9]{10})\s+([BS])\s+(NSE|BSE)\s+(.+)$"
    )
    number_pattern = r"\(?-?[\d,]+\.\d+\)?"
    rows, seen = [], set()
    for line_no, line in enumerate(lines, 1):
        match = trade_pattern.match(line)
        if not match:
            continue
        trade_id, symbol, isin, side, exchange, rest = match.groups()
        key = (exchange, trade_id)
        if key in seen:
            return [], ["Repeated annexure trade ID; review the source document."]
        seen.add(key)
        columns = re.split(r"\s{2,}", rest.strip())
        if len(columns) != 4 or not re.fullmatch(r"\d+", columns[0]):
            return [], [
                "Equity annexure has an unsupported rate or remarks column; review the document."
            ]
        units = int(columns[0])
        per_unit, net_rate = (amount(v) for v in columns[1:3])
        net_value = amount(columns[3].strip("()").lstrip("-"))
        # Zerodha reports brokerage separately: its net rate is the gross
        # execution rate, rather than a price to adjust for brokerage.
        price = net_rate
        if units <= 0 or price <= 0 or abs(net_rate * units - net_value) > Decimal("0.05"):
            return [], ["Equity annexure rates do not reconcile to its trade value."]
        rows.append(
            {
                "date": day,
                "symbol": canonical_symbol(symbol),
                "isin": isin,
                "side": "BUY" if side == "B" else "SELL",
                "units": units,
                "price": str(price),
                "trade_id": trade_id,
                "exchange": exchange,
                "components": {"brokerage": str(per_unit * units)},
                "group": "contract",
                "estimated": False,
                "line": line_no,
            }
        )
    if not rows:
        return [], ["No equity trades recognized in the Zerodha annexure."]
    # The ISIN summary is an independent completeness check. Daily fees must
    # never be spread over an annexure from which some equity trades were lost.
    expected, actual = {}, {}
    for line in lines:
        summary = re.match(r"^\s*(IN[A-Z0-9]{10})\s+\S+\s+(.+)$", line)
        if not summary:
            continue
        columns = re.split(r"\s{2,}", summary.group(2).strip())
        if len(columns) != 12:
            return [], ["Equity ISIN summary layout was not recognized."]
        for side, index in (("BUY", 0), ("SELL", 5)):
            if not re.fullmatch(r"\d+", columns[index]):
                return [], ["Equity ISIN summary quantity is invalid."]
            expected[(summary.group(1), side)] = int(columns[index])
    for row in rows:
        key = (row["isin"], row["side"])
        actual[key] = actual.get(key, 0) + row["units"]
    if (
        not expected
        or any(actual.get(key, 0) != units for key, units in expected.items())
        or any(key not in expected for key in actual)
    ):
        return [], [
            "Annexure quantities do not match the equity ISIN summary; no charges were allocated."
        ]
    header = next((line for line in lines if "NCL-Cash" in line and "NET TOTAL" in line), None)
    if not header or "NCL-F&O" not in header:
        return [], ["Zerodha cash charge columns were not recognized."]
    # Column boundary avoids the repeated NET TOTAL and derivative charges.
    cash_center = header.index("NCL-Cash") + len("NCL-Cash") / 2
    fo_center = header.index("NCL-F&O") + len("NCL-F&O") / 2
    boundary = int((cash_center + fo_center) / 2)

    def cash_value(line):
        numbers = re.findall(number_pattern, line[:boundary])
        if len(numbers) > 1:
            raise DomainValidationError("cash charge column has ambiguous amounts")
        if not numbers:
            return Decimal(0)
        value = numbers[0].replace(",", "")
        return -Decimal(value[1:-1]) if value.startswith("(") else Decimal(value)

    labels = {
        "brokerage": "Taxable value of Supply (Brokerage)",
        "exchange": "Exchange transaction charges",
        "clearing": "Clearing charges",
        "cgst": "CGST",
        "sgst": "SGST",
        "igst": "IGST",
        "stt": "Securities transaction tax",
        "sebi": "SEBI turnover fees",
        "stamp_duty": "Stamp duty",
    }
    totals, obligations = {}, {}
    for line in lines:
        stripped = line.strip()
        for component, label in labels.items():
            if stripped.startswith(label):
                value = cash_value(line)
                if value > 0:
                    return [], ["Charge refund or credit requires manual review."]
                if component in totals:
                    return [], ["Duplicate cash charge summary; review the document."]
                totals[component] = abs(value)
                break
        for key, label in (
            ("gross", "Pay in/Pay out obligation"),
            ("net", "Net amount receivable/(payable by client)"),
        ):
            if stripped.startswith(label):
                obligations[key] = cash_value(line)
    if set(totals) != set(labels) or set(obligations) != {"gross", "net"}:
        return [], ["Zerodha charge summary is incomplete; no fees were imported."]
    fee_total = sum(totals.values(), Decimal(0))
    if abs(obligations["gross"] - obligations["net"] - fee_total) > Decimal("0.01"):
        return [], ["Cash charges do not reconcile to the contract note's net obligation."]
    annexure_obligation = sum(
        (Decimal(r["price"]) * r["units"] * (1 if r["side"] == "SELL" else -1) for r in rows),
        Decimal(0),
    )
    if abs(annexure_obligation - obligations["gross"]) > Decimal("0.05"):
        return [], ["Equity annexure does not reconcile to the cash obligation."]
    totals["gst"] = sum((totals.pop(k) for k in ("cgst", "sgst", "igst")), Decimal(0))
    for component, total in totals.items():
        if component == "brokerage" and total == sum(
            (amount(r["components"]["brokerage"]) for r in rows), Decimal(0)
        ):
            continue
        weights = [
            Decimal(r["price"]) * r["units"]
            if component != "stamp_duty" or r["side"] == "BUY"
            else Decimal(0)
            for r in rows
        ]
        if total and not sum(weights):
            return [], ["Stamp duty has no corresponding purchase; review the document."]
        for row, share in zip(rows, allocate(total, weights), strict=True):
            row["components"][component] = str(share)
            row["estimated"] = row["estimated"] or (bool(total) and sum(w > 0 for w in weights) > 1)
    warnings = (
        [
            "Daily cash charges allocated across all equity trades, including outside trades; multi-trade allocations are estimates."
        ]
        if any(r["estimated"] for r in rows)
        else []
    )
    return rows, warnings


def parse_document(filename, raw, password="", kind="contract"):
    extension = Path(filename).suffix.lower()
    if not raw or len(raw) > 8 * 1024 * 1024:
        raise DomainValidationError("each document must be nonempty and no larger than 8 MB")
    if extension == ".pdf":
        return _parse_pdf(_pdf_text(raw, password), kind)
    if extension not in {".csv", ".xlsx"}:
        raise DomainValidationError("choose PDF, CSV or XLSX files")
    try:
        return _parse_tables(_tables(raw, extension), kind)
    except (BadZipFile, InvalidFileException, csv.Error) as exc:
        raise DomainValidationError("spreadsheet or CSV could not be read") from exc
